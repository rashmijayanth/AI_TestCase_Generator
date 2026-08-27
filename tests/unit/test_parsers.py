import io

import pytest
from docx import Document as DocxDocument

from testgen.ingestion import parsers
from testgen.platform.enums import DocumentFormat


@pytest.mark.unit
def test_detect_format_by_extension() -> None:
    assert parsers.detect_format("srs.pdf") == DocumentFormat.PDF
    assert parsers.detect_format("srs.docx") == DocumentFormat.DOCX
    assert parsers.detect_format("srs.xml") == DocumentFormat.XML
    assert parsers.detect_format("srs.reqif") == DocumentFormat.REQIF
    assert parsers.detect_format("srs.md") == DocumentFormat.MARKDOWN


@pytest.mark.unit
def test_detect_format_rejects_unknown_extension() -> None:
    with pytest.raises(ValueError, match="Unsupported"):
        parsers.detect_format("srs.txt")


@pytest.mark.unit
def test_parse_docx_extracts_paragraph_text() -> None:
    document = DocxDocument()
    document.add_paragraph("REQ-001: The system shall log all alarms.")
    document.add_paragraph("")  # blank paragraphs should be dropped
    document.add_paragraph("REQ-002: The system shall stop within 500ms.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = parsers.parse_docx(buffer.getvalue())

    assert "REQ-001: The system shall log all alarms." in text
    assert "REQ-002: The system shall stop within 500ms." in text


@pytest.mark.unit
def test_parse_markdown_strips_markup() -> None:
    raw = b"# Requirements\n\nREQ-001: The system **shall** log all alarms.\n"

    text = parsers.parse_markdown(raw)

    assert "REQ-001: The system" in text
    assert "shall" in text
    assert "**" not in text


@pytest.mark.unit
def test_parse_xml_extracts_text_content() -> None:
    raw = b"<requirements><req id='1'>The system shall beep.</req></requirements>"

    text = parsers.parse_xml(raw)

    assert text == "The system shall beep."


@pytest.mark.unit
def test_parse_reqif_extracts_spec_object_values() -> None:
    raw = b"""<?xml version="1.0" encoding="UTF-8"?>
    <REQ-IF xmlns="http://www.omg.org/spec/ReqIF/20110401/reqif.xsd">
      <CORE-CONTENT>
        <REQ-IF-CONTENT>
          <SPEC-OBJECTS>
            <SPEC-OBJECT IDENTIFIER="sobj-1">
              <VALUES>
                <ATTRIBUTE-VALUE-STRING>
                  <THE-VALUE>The system shall beep on occlusion.</THE-VALUE>
                </ATTRIBUTE-VALUE-STRING>
              </VALUES>
            </SPEC-OBJECT>
          </SPEC-OBJECTS>
        </REQ-IF-CONTENT>
      </CORE-CONTENT>
    </REQ-IF>"""

    text = parsers.parse_reqif(raw)

    assert text == "The system shall beep on occlusion."


@pytest.mark.unit
def test_parse_pdf_joins_page_text(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        def __init__(self, _stream: io.BytesIO) -> None:
            self.pages = [FakePage("Page one text."), FakePage(""), FakePage("Page two text.")]

    monkeypatch.setattr(parsers, "PdfReader", FakeReader)

    text = parsers.parse_pdf(b"irrelevant-bytes")

    assert text == "Page one text.\n\nPage two text."


@pytest.mark.unit
def test_extract_text_dispatches_on_format() -> None:
    text = parsers.extract_text(DocumentFormat.MARKDOWN, b"REQ-1: The system shall log alarms.\n")

    assert "shall log alarms" in text
